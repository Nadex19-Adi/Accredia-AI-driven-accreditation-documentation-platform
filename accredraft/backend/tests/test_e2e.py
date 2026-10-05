"""
End-to-end workflow test against a running API.

Covers: login -> template -> project -> upload -> ingest -> extract ->
review edit -> validate -> generate -> download -> audit.

Run:  .venv\\Scripts\\python.exe tests\\test_e2e.py
Requires the API to be running (scripts\\start_server.py).
"""
import io
import json
import sys
import time
import urllib.error
import urllib.request

import docx

BASE = "http://127.0.0.1:8000"


def req(method, path, token=None, json_body=None, data=None, headers=None):
    hdrs = dict(headers or {})
    if token:
        hdrs["Authorization"] = f"Bearer {token}"
    body = None
    if json_body is not None:
        body = json.dumps(json_body).encode()
        hdrs["Content-Type"] = "application/json"
    elif data is not None:
        body = data
    request = urllib.request.Request(BASE + path, data=body, headers=hdrs, method=method)
    try:
        with urllib.request.urlopen(request, timeout=60) as resp:
            raw = resp.read()
            ctype = resp.headers.get("Content-Type", "")
            return resp.status, (json.loads(raw) if "application/json" in ctype else raw)
    except urllib.error.HTTPError as exc:
        return exc.code, exc.read().decode(errors="replace")


def make_sample_docx() -> bytes:
    doc = docx.Document()
    doc.add_heading("Event Report", 0)
    for line in [
        "Title: Workshop on AI Ethics",
        "Date: 15-12-2024",
        "Venue: Main Auditorium",
        "Organizer: Department of Computer Science",
        "Participants: 45",
    ]:
        doc.add_paragraph(line)
    doc.add_paragraph("Objectives:")
    doc.add_paragraph("1. Understand core principles of AI ethics and fairness")
    doc.add_paragraph("2. Apply ethical frameworks to real deployment scenarios")
    buffer = io.BytesIO()
    doc.save(buffer)
    return buffer.getvalue()


def multipart(files):
    boundary = "----accredraft" + str(int(time.time() * 1000))
    parts = []
    for name, (filename, content, ctype) in files.items():
        parts.append(
            f"--{boundary}\r\nContent-Disposition: form-data; name=\"{name}\"; "
            f"filename=\"{filename}\"\r\nContent-Type: {ctype}\r\n\r\n".encode()
            + content + b"\r\n"
        )
    parts.append(f"--{boundary}--\r\n".encode())
    return b"".join(parts), f"multipart/form-data; boundary={boundary}"


def main() -> int:
    checks = []

    def check(name, ok, detail=""):
        checks.append((name, ok, detail))
        print(f"  [{'PASS' if ok else 'FAIL'}] {name}" + (f" :: {detail}" if detail and not ok else ""))

    print("\n1. Auth")
    status, body = req("POST", "/api/auth/login",
                       json_body={"email": "admin@example.com", "password": "admin12345"})
    if status != 200:
        print("   login failed:", status, body)
        return 1
    token = body["access_token"]
    check("login", bool(token))

    print("\n2. Templates")
    status, templates = req("GET", "/api/templates", token=token)
    check("list templates", status == 200 and len(templates) > 0, f"{status} {templates}")
    if not templates:
        return 1
    template = templates[0]
    print(f"   using: {template['body']} / {template['name']} v{template['version']} ({template['status']})")

    print("\n3. Project")
    status, project = req("POST", "/api/projects", token=token, json_body={
        "template_id": template["id"], "name": "E2E Test Event Report",
        "description": "Automated end-to-end test",
    })
    check("create project", status == 200, f"{status} {project}")
    if status != 200:
        return 1
    pid = project["id"]
    check("template version pinned", project["template_version"] == template["version"])

    print("\n4. Upload")
    body_bytes, ctype = multipart({"files": (
        "event_report.docx", make_sample_docx(),
        "application/vnd.openxmlformats-officedocument.wordprocessingml.document")})
    status, files = req("POST", f"/api/projects/{pid}/files", token=token, data=body_bytes,
                        headers={"Content-Type": ctype})
    check("upload file", status == 200 and len(files) == 1, f"{status} {files}")
    if status != 200:
        return 1
    fid = files[0]["id"]
    check("file kind detected", files[0]["file_kind"] == "docx", files[0]["file_kind"])

    print("\n5. Ingest")
    status, _ = req("POST", f"/api/projects/{pid}/process", token=token)
    check("trigger process", status == 200, str(status))
    evidence = []
    for _ in range(30):
        time.sleep(1)
        status, evidence = req("GET", f"/api/files/{fid}/evidence", token=token)
        if status == 200 and evidence:
            break
    check("evidence items parsed", len(evidence) > 0, f"count={len(evidence)}")

    print("\n6. Extract")
    status, _ = req("POST", f"/api/projects/{pid}/extract", token=token)
    check("trigger extract", status == 200, str(status))
    fields = []
    for _ in range(30):
        time.sleep(1)
        status, fields = req("GET", f"/api/projects/{pid}/extractions", token=token)
        if status == 200 and any(f["status"] != "missing" for f in fields):
            break
    by_id = {f["field_id"]: f for f in fields}
    print("   extracted values:")
    for f in fields:
        print(f"     - {f['field_id']}: {json.dumps(f['value_json'])[:70]} ({f['status']})")
    check("fields extracted", len(fields) > 0)
    check("title extracted", by_id.get("title", {}).get("value_json") is not None)
    check("date extracted", by_id.get("date", {}).get("value_json") is not None)
    check("participant_count extracted", by_id.get("participant_count", {}).get("value_json") is not None)
    check("objectives extracted as list",
          isinstance(by_id.get("objectives", {}).get("value_json"), list))
    check("unfound field stays null", by_id.get("photos", {}).get("value_json") is None)

    print("\n7. Validate (before review)")
    status, validation = req("POST", f"/api/projects/{pid}/validate/sync", token=token)
    check("run validation", status == 200, f"{status} {validation}")
    if status == 200:
        check("required-field errors reported", validation.get("errors", 0) > 0, str(validation))

    print("\n8. Human review")
    title_field = by_id.get("title")
    if title_field:
        status, _ = req("PUT", f"/api/projects/{pid}/extractions/{title_field['id']}", token=token,
                        json_body={"value": "Workshop on AI Ethics (reviewed)", "status": "edited"})
        check("edit a field", status == 200, str(status))
    attendance = by_id.get("attendance")
    if attendance:
        status, _ = req("PUT", f"/api/projects/{pid}/extractions/{attendance['id']}", token=token,
                        json_body={"status": "na"})
        check("mark field N/A", status == 200, str(status))

    status, validation = req("POST", f"/api/projects/{pid}/validate/sync", token=token)
    if status == 200:
        print(f"   after review: ok={validation['ok']} errors={validation['errors']} "
              f"warnings={validation['warnings']}")

    print("\n9. Generate")
    status, gen = req("POST", f"/api/projects/{pid}/generate/sync", token=token)
    check("generate document", status == 200 and "docx_key" in gen, f"{status} {gen}")

    if status == 200:
        print("\n10. Download")
        status, outputs = req("GET", f"/api/projects/{pid}/outputs", token=token)
        check("list outputs", status == 200 and len(outputs) > 0, str(status))
        docx_out = next((o for o in outputs if o["format"] == "docx"), None)
        if docx_out:
            status, blob = req("GET", f"/api/outputs/{docx_out['id']}/download", token=token)
            ok = status == 200 and isinstance(blob, bytes) and blob[:2] == b"PK"
            check("download docx (valid zip)", ok, f"status={status}")
            print(f"   docx size: {len(blob)} bytes")

    print("\n11. Audit")
    status, audit = req("GET", f"/api/projects/{pid}/audit", token=token)
    check("audit log populated", status == 200 and len(audit) > 0,
          f"entries={len(audit) if status == 200 else status}")

    # Every endpoint the web client calls must exist with a matching path.
    print("\n12. API surface (frontend contract)")
    status, me = req("GET", "/api/auth/me", token=token)
    check("GET /auth/me", status == 200 and me.get("email") == "admin@example.com", str(status))

    status, one = req("GET", f"/api/templates/{template['id']}", token=token)
    check("GET /templates/{id}", status == 200 and one["id"] == template["id"], str(status))

    status, run = req("GET", f"/api/projects/{pid}/validation", token=token)
    check("GET /projects/{id}/validation", status == 200 and len(run.get("results", [])) > 0,
          str(status))

    status, docs = req("GET", f"/api/projects/{pid}/documents", token=token)
    check("GET /projects/{id}/documents", status == 200 and len(docs) > 0, str(status))
    if status == 200 and docs:
        check("document model versioned", docs[0]["version"] >= 1, str(docs[0]["version"]))

    status, _ = req("DELETE", f"/api/files/{fid}", token=token)
    check("DELETE /files/{id}", status == 200, str(status))

    status, _ = req("DELETE", f"/api/projects/{pid}", token=token)
    check("DELETE /projects/{id}", status == 200, str(status))

    status, _ = req("GET", f"/api/projects/{pid}", token=token)
    check("deleted project is gone", status == 404, str(status))

    passed = sum(1 for _, ok, _ in checks if ok)
    print(f"\n{'='*58}\nRESULT: {passed}/{len(checks)} checks passed\n{'='*58}")
    for name, ok, detail in checks:
        if not ok:
            print(f"  FAILED: {name} {detail}")
    return 0 if passed == len(checks) else 1


if __name__ == "__main__":
    raise SystemExit(main())
